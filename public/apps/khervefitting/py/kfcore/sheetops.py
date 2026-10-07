# KherveFitting - XPS Data Analysis Software
# Copyright (C) 2024-2026 Gwilherm Kerherve <g.kerherve@ic.ac.uk>
#
# KherveFitting is dual-licensed:
#   - GNU GPL v3.0 (see LICENSE-GPL.txt) for open-source use
#   - Commercial Licence (see LICENSE-COMMERCIAL.txt) for proprietary use
# SPDX-License-Identifier: GPL-3.0-only OR LicenseRef-KherveFitting-Commercial
#
# KherveOS port: the core-level (sheet) operations of the Edit menu and the
# toolbar — Utilities.rename_sheet, on_delete_sheet, copy_sheet, CropWindow,
# JoinSheetsWindow and sort_excel_sheets.
#
# The desktop edits the open .xlsx at once. Here the workbook is the bytes the
# page opened (``original``, may be None for an imported file); each operation
# returns the new sheet order and the edited bytes, so a later save writes the
# renamed / deleted / new / reordered sheets where the desktop has them.

import copy
import io
import re

import numpy as np

SPECIAL_SHEETS = ("Results Table", "Experimental Description")


class SheetError(ValueError):
    """The desktop refuses (its message box text)."""


# ── The workbook bytes ──────────────────────────────────────────────────
def _load(original):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(original))


def _dump(wb):
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _edit(original, fn):
    """Apply fn(workbook) to the bytes (nothing to do without a workbook)."""
    if not original:
        return original
    wb = _load(original)
    fn(wb)
    return _dump(wb)


def _write_sheet(wb, name, columns):
    """pd.DataFrame(columns).to_excel(sheet_name=name, index=False, mode='a', if_sheet_exists='replace')."""
    if name in wb.sheetnames:
        idx = wb.sheetnames.index(name)
        wb.remove(wb[name])
        ws = wb.create_sheet(name, idx)
    else:
        ws = wb.create_sheet(name)
    for c, (header, values) in enumerate(columns, start=1):
        ws.cell(row=1, column=c, value=header)
        for r, v in enumerate(values, start=2):
            ws.cell(row=r, column=c, value=None if v is None else float(v))


def _row_of(sheet_name):
    m = re.search(r'(\d+)$', sheet_name)
    return m.group(1) if m else "0"


# ── Rename ──────────────────────────────────────────────────────────────
def rename_sheet(w, sheets, original, old, new):
    """Utilities.rename_sheet (and the Sample Manager's duplicate check)."""
    from .sheets import select_sheet
    new = (new or '').strip()
    if not new or new == old:
        return sheets, original, old
    if len(new.split()) > 1:
        raise SheetError("Only one single word is allowed for core levels or sheet names.")
    if old not in w.Data['Core levels']:
        raise SheetError(f"No core level named '{old}'.")
    if new in w.Data['Core levels']:
        raise SheetError(f"A core level named '{new}' already exists.")
    if len(new) > 31:
        raise SheetError("Excel sheet names are limited to 31 characters.")

    def fn(wb):
        if old in wb.sheetnames:
            wb[old].title = new
    original = _edit(original, fn)
    # The core level keeps its place in Data (the desktop's pop/insert moves it last;
    # the order shown is the workbook's, which keeps the sheet where it was).
    w.Data['Core levels'] = {(new if k == old else k): v for k, v in w.Data['Core levels'].items()}
    sheets = [new if s == old else s for s in sheets]
    select_sheet(w, new)
    return sheets, original, new


# ── Delete ──────────────────────────────────────────────────────────────
def delete_sheet(w, sheets, original, name):
    """Utilities.on_delete_sheet: the first remaining sheet is selected."""
    from .sheets import select_sheet

    def fn(wb):
        if name in wb.sheetnames:
            wb.remove(wb[name])
    original = _edit(original, fn)
    if name in w.Data['Core levels']:
        del w.Data['Core levels'][name]
        w.Data['Number of Core levels'] = max(0, int(w.Data.get('Number of Core levels', 1)) - 1)
    sheets = [s for s in sheets if s != name]
    selected = sheets[0] if sheets else ''
    if selected:
        select_sheet(w, selected)
    else:
        select_sheet(w, '')
    return sheets, original, selected


# ── Copy (Edit > Copy Core Level = Utilities.copy_sheet) ────────────────
def copy_name(w, sheet_name):
    """The name copy_sheet gives the copy: the next free number of its core level."""
    match = re.match(r'([A-Za-z]+(?:\d+[spdfg]+)?)(\d*)$', sheet_name)
    base_name = match.group(1) if match else sheet_name
    existing = list(w.Data['Core levels'].keys())
    max_suffix = 0
    for s in existing:
        if not re.match(r'^' + re.escape(base_name) + r'\d*$', s):
            continue
        m = re.search(r'^' + re.escape(base_name) + r'(\d+)$', s)
        if m:
            max_suffix = max(max_suffix, int(m.group(1)))
    if base_name in existing or f"{base_name}0" in existing or max_suffix > 0:
        new = f"{base_name}{max_suffix + 1}"
    else:
        new = base_name
    while new in existing:
        m = re.search(r'^' + re.escape(base_name) + r'(\d+)$', new)
        new = f"{base_name}{int(m.group(1)) + 1}" if m else f"{base_name}1"
    return new


def copy_sheet(w, sheets, original, sheet_name):
    """Utilities.copy_sheet: a copy of the core level, appended and selected."""
    from .sheets import select_sheet
    if not sheet_name or sheet_name not in w.Data['Core levels']:
        raise SheetError("No data present in the plot to be copied.\nPlease select a dataset first.")
    new = copy_name(w, sheet_name)
    src = w.Data['Core levels'][sheet_name]
    # The desktop's dict.copy() shares Fitting/Background with the source until the
    # file is reopened; a deep copy is what the reopened file holds.
    w.Data['Core levels'][new] = copy.deepcopy(src)
    w.Data['Number of Core levels'] = int(w.Data.get('Number of Core levels', 0)) + 1
    be = src['B.E.']
    bkg = (src.get('Background') or {}).get('Bkg Y') or src['Raw Data']
    original = _edit(original, lambda wb: _write_sheet(wb, new, [
        ('BE', be), ('Raw Data', src['Raw Data']), ('Background', bkg), ('Transmission', [1.0] * len(be))]))
    sheets = list(sheets) + [new]
    select_sheet(w, new)
    return sheets, original, new


# ── Crop ────────────────────────────────────────────────────────────────
def _crop_base(sheet_name):
    if '~' in sheet_name:
        return re.sub(r'\d+$', '', sheet_name)
    m = re.match(r'([A-Za-z]+\d*[spdfg]*)', sheet_name)
    return m.group(1) if m else sheet_name


def crop_info(w, sheet_name):
    """CropWindow.init_values: default range, suggested name and the sheets of the same core level."""
    if not sheet_name or sheet_name not in w.Data['Core levels']:
        raise SheetError("No data present in the plot to be cropped.\nPlease select a dataset first.")
    x = w.Data['Core levels'][sheet_name]['B.E.']
    base = _crop_base(sheet_name)
    names = list(w.Data['Core levels'].keys())
    # get_earliest_row_name
    if base not in names:
        suggested = base
    else:
        used = set()
        pattern = re.compile(f"^{re.escape(base)}(\\d+)$")
        for s in names:
            if s == base:
                used.add(0)
            else:
                m = pattern.match(s)
                if m:
                    used.add(int(m.group(1)))
        suggested = next((base if i == 0 else f"{base}{i}" for i in range(1000) if i not in used),
                         f"{base}{len(names)}")
    pattern = re.compile(f"^{re.escape(base)}\\d*$")

    def natural(s):
        m = re.search(r'\d+$', s)
        return (s[:m.start()], int(m.group())) if m else (s, 0)
    related = sorted([s for s in names if pattern.match(s)], key=natural)
    return {'low': round(min(x) + 2, 2), 'high': round(max(x) - 2, 2), 'name': suggested, 'related': related}


def _crop_name(w, sheet_to_crop):
    """CropWindow.on_crop: the first free number of the sheet's core level."""
    base = _crop_base(sheet_to_crop)
    existing = list(w.Data['Core levels'].keys())
    used = set()
    for s in existing:
        m = re.search(r'^' + re.escape(base) + r'(\d+)$', s)
        if m:
            used.add(int(m.group(1)))
        elif s == base:
            used.add(0)
    suffix = 0
    while suffix in used:
        suffix += 1
    return base if (suffix == 0 and base not in existing) else f"{base}{suffix}"


def crop_sheets(w, sheets, original, sheet_names, low, high, name=None, name_sheet=None):
    """CropWindow.on_crop for the ticked sheets. Returns (sheets, original, created, failures).

    sheet_names are cropped in the order of the window's list (natural sort).
    name: the window's "New name" box. The desktop computes each new name itself
    (its box only shows the suggestion); a name typed here is used for name_sheet
    (default: the first sheet) when it is free.
    """
    from .sheets import select_sheet
    min_be, max_be = float(min(low, high)), float(max(low, high))
    created, failures = [], []
    new_columns = {}

    def natural(s):
        m = re.search(r'\d+$', s)
        return (s[:m.start()], int(m.group())) if m else (s, 0)
    sheet_names = sorted(dict.fromkeys(sheet_names), key=natural)
    if name_sheet is None and sheet_names:
        name_sheet = sheet_names[0]
    for sheet in sheet_names:
        try:
            if sheet not in w.Data['Core levels']:
                raise SheetError("no such sheet")
            new = _crop_name(w, sheet)
            if sheet == name_sheet and name and name.strip() and name.strip() not in w.Data['Core levels']:
                if len(name.split()) > 1:
                    raise SheetError("Only one single word is allowed for core levels or sheet names.")
                new = name.strip()
            data = w.Data['Core levels'][sheet]
            x = np.array(data['B.E.'], dtype=float)
            mask = (x >= min_be) & (x <= max_be)
            if not mask.any():
                raise SheetError("no data points in this range")
            bkg = (data.get('Background') or {}).get('Bkg Y') or data['Raw Data']
            new_data = {
                'B.E.': x[mask].tolist(),
                'Raw Data': np.array(data['Raw Data'], dtype=float)[mask].tolist(),
                'Background': {'Bkg Y': np.array(bkg, dtype=float)[mask].tolist()},
                'Name': new,
            }
            w.Data['Core levels'][new] = new_data
            w.Data['Number of Core levels'] = int(w.Data.get('Number of Core levels', 0)) + 1
            new_columns[new] = [
                ('BE', [float(f"{v:.2f}") for v in new_data['B.E.']]),
                ('Raw Data', [float(f"{v:.2f}") for v in new_data['Raw Data']]),
                ('Background', [float(f"{v:.2f}") for v in new_data['Background']['Bkg Y']]),
                ('Transmission', [1.00] * int(mask.sum())),
            ]
            created.append(new)
        except Exception as e:  # the desktop lists failures and goes on
            failures.append(f"{sheet}: {e}")

    def fn(wb):
        for new, cols in new_columns.items():
            _write_sheet(wb, new, cols)
    original = _edit(original, fn) if new_columns else original
    sheets = list(sheets) + [n for n in created if n not in sheets]
    if created:
        select_sheet(w, created[0])
    return sheets, original, created, failures


# ── Join ────────────────────────────────────────────────────────────────
def join_sheets(w, sheets, original, selected, name="Joined_Scan"):
    """JoinSheetsWindow.on_join: spectra end to end, the highest BE first."""
    from .sheets import select_sheet
    selected = [s for s in selected if s in w.Data['Core levels']]
    if len(selected) < 2:
        raise SheetError("Select at least 2 sheets to join")
    name = (name or "Joined_Scan").strip()
    if len(name.split()) > 1:
        raise SheetError("Only one single word is allowed for core levels or sheet names.")
    if name in w.Data['Core levels']:
        # pandas' ExcelWriter(mode='a') refuses an existing sheet name.
        raise SheetError(f"Sheet '{name}' already exists.")
    max_be = {s: max(w.Data['Core levels'][s]['B.E.']) for s in selected}
    ordered = sorted(selected, key=lambda s: max_be[s], reverse=True)
    be, data = [], []
    for s in ordered:
        be.extend(w.Data['Core levels'][s]['B.E.'])
        data.extend(w.Data['Core levels'][s]['Raw Data'])
    w.Data['Core levels'][name] = {
        'B.E.': be,
        'Raw Data': data,
        'Background': {'Bkg Y': list(data)},
        'Transmission': [1.0] * len(be),
    }
    w.Data['Number of Core levels'] = len(w.Data['Core levels'])
    original = _edit(original, lambda wb: _write_sheet(wb, name, [
        ('BE', be), ('Raw Data', data), ('Background', data), ('Transmission', [1.0] * len(be))]))
    sheets = list(sheets) + [name]
    select_sheet(w, name)
    return sheets, original, name


# ── Sort ────────────────────────────────────────────────────────────────
def sorted_names(sheet_names):
    """sort_excel_sheets: by sample number, then core level (Fermi, VB, Survey last)."""
    regular, results, expdesc = [], None, None
    for s in sheet_names:
        if s == "Results Table":
            results = s
        elif s == "Experimental Description":
            expdesc = s
        else:
            regular.append(s)
    groups = {}
    for s in regular:
        low = s.lower()
        if "wide" in low or "survey" in low:
            m = re.match(r'(wide|survey)(\d*)$', low, re.IGNORECASE)
            base, num = (m.group(1).capitalize(), m.group(2)) if m else (s, "")
        elif low.startswith('vb'):
            m = re.match(r'(vb)(\d*)$', s, re.IGNORECASE)
            base, num = m.groups() if m else (s, "")
        elif low.startswith('valence'):
            m = re.match(r'(valence)(\d*)$', s, re.IGNORECASE)
            base, num = m.groups() if m else (s, "")
        elif low.startswith('fermi'):
            m = re.match(r'(fermi)(\d*)$', s, re.IGNORECASE)
            base, num = m.groups() if m else (s, "")
        else:
            m = re.match(r'([A-Za-z]+\d*[spdfg]*)(\d*)$', s)
            base, num = m.groups() if m else (s, "")
        groups.setdefault(int(num) if num else 0, []).append((base, s))

    def key(item):
        base = item[0].lower()
        if "wide" in base or "survey" in base:
            return "zzz"
        if base in ("vb", "valence"):
            return "yy"
        if base == "fermi":
            return "yx"
        return base
    out = []
    for num in sorted(groups):
        out.extend(s for _, s in sorted(groups[num], key=key))
    if results:
        out.append(results)
    if expdesc:
        out.append(expdesc)
    return out


def sort_sheets(w, sheets, original):
    """sort_excel_sheets: workbook, Data and the sheet list in the same order."""
    order = sorted_names(list(sheets))

    def fn(wb):
        full = sorted_names(wb.sheetnames)
        for i, s in enumerate(full):
            ws = wb[s]
            wb.move_sheet(ws, offset=i - wb.index(ws))
    original = _edit(original, fn)
    cls = w.Data['Core levels']
    rank = {s: i for i, s in enumerate(sorted_names(list(cls.keys())))}
    w.Data['Core levels'] = {k: cls[k] for k in sorted(cls, key=lambda k: rank[k])}
    current = w.sheet_combobox.GetValue()
    if current not in order and order:
        from .sheets import select_sheet
        select_sheet(w, order[0])
        current = order[0]
    return order, original, current
