# KherveFitting - XPS Data Analysis Software
# Copyright (C) 2024-2026 Gwilherm Kerherve <g.kerherve@ic.ac.uk>
#
# KherveFitting is dual-licensed:
#   - GNU GPL v3.0 (see LICENSE-GPL.txt) for open-source use
#   - Commercial Licence (see LICENSE-COMMERCIAL.txt) for proprietary use
# SPDX-License-Identifier: GPL-3.0-only OR LicenseRef-KherveFitting-Commercial
#
# KherveOS port: KherveFitting's own file format. A workbook (.xlsx) has one
# sheet per core level ("Binding Energy", "Raw Data" in row 1) and, next to
# it, a .json file with everything else (window.Data: backgrounds, peaks,
# constraints, results tables). Opening follows FileMenu/Open.open_xlsx_file
# and ConfigFile.add_core_level_Data; saving follows FileMenu/Save.save_data,
# save_to_excel and save_results_table. Synced with KherveFitting-AI (dev-AI
# v1.93): small numbers keep significant figures in the JSON, and the results
# sheet has the weight % and the seven 1-sigma uncertainty columns.

import io
import json
import re

import numpy as np

from .curves import overall_fit, peak_curve
from .grid import PEAK_COLUMNS
from .session import empty_data
from .uncertainty import RESULT_ERROR_COLUMNS, RESULT_ERROR_FIRST_COL

RESULT_COLUMNS = [
    "Peak\nLabel", "Position\n(eV)", "Height\n(CPS)", "FWHM\n(eV)", "L/G \nσ/γ (%)",
    "Area\n(CPS.eV)", "Atomic\n(%)", " ", "RSF", "TXFN", "ECF", "Instr.", "Fitting Model",
    "Corr. Area\n(a.u.)", "σ or α\nW_g", "γ or β\nW_l", "Bkg Type", "Bkg Low\n(eV)",
    "Bkg High\n(eV)", "Bkg Offset Low\n(CPS)", "Bkg Offset High\n(CPS)", "Sheetname", "Position\nConstraint",
    "Height\nConstraint", "FWHM\nConstraint", "L/G\nConstraint", "Area\nConstraint", "σ\nConstraint",
    "γ\nConstraint", "Weight\n(%)", "Mass\n(amu)"] + [hdr for _key, _qty, hdr in RESULT_ERROR_COLUMNS]

RESULT_FIELDS = ['Name', 'Position', 'Height', 'FWHM', 'L/G', 'Area', 'at. %', 'Checkbox', 'RSF', 'TXFN', 'ECF',
                 'Instrument', 'Fitting Model', 'Rel. Area', 'Sigma', 'Gamma', 'Bkg Type', 'Bkg Low', 'Bkg High',
                 'Bkg Offset Low', 'Bkg Offset High', 'Sheetname', 'Pos. Constraint', 'Height Constraint',
                 'FWHM Constraint', 'L/G Constraint', 'Area Constraint', 'Sigma Constraint', 'Gamma Constraint',
                 'wt. %', None]
RESULT_FIELDS += [None] * (RESULT_ERROR_FIRST_COL - len(RESULT_FIELDS)) + [key for key, _q, _h in RESULT_ERROR_COLUMNS]

#: Kinds of sheets the desktop opens elsewhere (maps, profiles, EDX/EELS).
SPECIAL = re.compile(r'^(zzProfile|zzMap|XPS~Map|EDX~|EELS~)|~Map')


def _openpyxl():
    import openpyxl
    return openpyxl


def _is_valid_header(c1, c2):
    """The header test of open_xlsx_file (XPS, Raman, XAS, EDX, EELS)."""
    col1 = str(c1).strip().upper()
    col2 = str(c2).strip().upper()
    xps = ('BE' in col1 or 'B.E.' in col1 or 'BINDING' in col1) and \
          ('RAW DATA' in col2 or 'CORRECTED DATA' in col2 or 'INTENSITY' in col2)
    raman = ('WAVENUMBER' in col1 or 'CM-1' in col1) and ('RAW DATA' in col2 or 'INTENSITY' in col2)
    xas = ('ENERGY' in col1 or 'PHOTON ENERGY' in col1) and ('INTENSITY' in col2 or 'RAW DATA' in col2)
    edx = ('ENERGY' in col1 and 'KEV' in col1) and 'INTENSITY' in col2
    eels = 'ENERGY' in col1 and ('EV' in col1 or 'LOSS' in col1) and 'INTENSITY' in col2
    return xps or raman or xas or edx or eels


def _num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if v != v else float(v)
    try:
        return float(str(v).strip())
    except ValueError:
        return None


def core_level_from_rows(name, rows, exp_info=None):
    """ConfigFile.add_core_level_Data for an ordinary spectrum sheet."""
    be_values, raw_data, corrected_data, transmission = [], [], [], []
    for r in rows[1:]:
        be = _num(r[0]) if len(r) > 0 else None
        raw = _num(r[1]) if len(r) > 1 else None
        if be is None or raw is None:
            continue
        be_values.append(be)
        raw_data.append(raw)
        c = _num(r[2]) if len(r) > 2 else None
        corrected_data.append(c if c is not None else raw)
        t = _num(r[3]) if len(r) > 3 else None
        transmission.append(t if t is not None else 1.0)
    cl = {
        'B.E.': [float(f"{v:.2f}") for v in be_values],
        'Raw Data': [float(f"{v:.2f}") for v in raw_data],
        'Corrected Data': [float(f"{v:.2f}") for v in corrected_data],
        'Transmission': [float(f"{v:.2f}") for v in transmission],
        'Name': name,
    }
    if exp_info:
        cl['ExperimentalInfo'] = exp_info
    if be_values:
        cl['Background'] = {
            'Bkg Y': cl['Raw Data'],
            'Bkg Type': '',
            'Bkg Low': float(f"{min(be_values):.2f}"),
            'Bkg High': float(f"{max(be_values):.2f}"),
            'Bkg Offset Low': 0,
            'Bkg Offset High': 0,
        }
    return cl


def _experimental_info(ws_rows):
    """The "Experimental Description" columns (40..60) of a sheet, if any."""
    if not ws_rows:
        return {}
    header = ws_rows[0]
    exp_col = None
    for col in range(39, min(60, len(header))):
        if header[col] is not None and "Experimental Description" in str(header[col]):
            exp_col = col
            break
    info = {}
    if exp_col is None:
        return info
    for r in ws_rows[1:]:
        if len(r) > exp_col and r[exp_col] is not None and str(r[exp_col]).strip():
            v = r[exp_col + 1] if len(r) > exp_col + 1 else None
            info[str(r[exp_col]).strip()] = str(v).strip() if v is not None else ""
    return info


def open_workbook(xlsx_bytes, json_text=None, file_path=''):
    """Read a KherveFitting workbook. Returns (Data, sheet names, dismissed).

    dismissed lists (sheet, reason) for sheets the desktop would not open.
    """
    openpyxl = _openpyxl()
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    names = [n for n in wb.sheetnames if n.lower() not in ("results table", "experimental description")]
    invalid = [n for n in names if re.match(r'^Sheet\d+$', n, re.IGNORECASE)]
    dismissed = [(n, "default Excel name (e.g. Sheet1)") for n in invalid]
    names = [n for n in names if n not in invalid]
    from .sheets import is_xps_like_sheet

    rows_by_sheet = {}
    valid = []
    for name in names:
        ws = wb[name]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if SPECIAL.search(name):
            dismissed.append((name, "maps, profiles and EDX/EELS sheets are not supported here yet"))
            continue
        if len(rows) == 0 or max((len(r) for r in rows), default=0) < 2:
            dismissed.append((name, "fewer than 2 columns or no data"))
            continue
        # dev-AI: every non-XPS technique writes its own axis label into
        # column A, so the sheet name (its technique prefix) decides for those.
        if not (_is_valid_header(rows[0][0], rows[0][1]) or not is_xps_like_sheet(None, name)):
            dismissed.append((name, f"unrecognised column headers (Col1='{rows[0][0]}', Col2='{rows[0][1]}')"))
            continue
        rows_by_sheet[name] = rows
        valid.append(name)

    if json_text:
        data = json.loads(json_text)
        data.setdefault('Core levels', {})
    else:
        data = empty_data()
    data['FilePath'] = file_path
    for name in valid:
        cl = data['Core levels'].get(name)
        if not cl or 'B.E.' not in cl or 'Raw Data' not in cl:
            data['Core levels'][name] = core_level_from_rows(name, rows_by_sheet[name],
                                                             _experimental_info(rows_by_sheet[name]))
    data['Number of Core levels'] = len(data['Core levels'])
    return data, valid, dismissed


# ── Saving ───────────────────────────────────────────────────────────────
def _round_keep_small(value, decimal_places):
    """Save._round_keep_small: round(value, decimal_places), except that a
    number below 1 keeps five significant figures instead of being rounded
    to 0.00 - small peak areas, their uncertainties, normalised intensities."""
    value = float(value)
    a = abs(value)
    if a == 0 or a >= 1 or not np.isfinite(value):
        return round(value, decimal_places) if np.isfinite(value) else value
    return round(value, max(decimal_places, 4 - int(np.floor(np.log10(a)))))


def to_json_data(obj, decimal_places=2):
    """Save.convert_to_serializable_and_round: floats rounded to 2 decimals
    (5 significant figures below 1)."""
    if isinstance(obj, bool):
        return obj
    if isinstance(obj, (float, np.floating)):
        return _round_keep_small(obj, decimal_places)
    if isinstance(obj, (int, np.integer)):
        return int(obj)
    if isinstance(obj, np.ndarray):
        return [to_json_data(v, decimal_places) for v in obj.tolist()]
    if isinstance(obj, (list, tuple)):
        return [to_json_data(v, decimal_places) for v in obj]
    if isinstance(obj, dict):
        return {k: to_json_data(v, decimal_places) for k, v in obj.items()}
    return obj


def _nan_to_none(values):
    out = []
    for v in values:
        f = float(v)
        out.append(None if f != f else f)
    return out


def _write_fit_columns(ws, window, name):
    """save_to_excel: columns F onwards (fit curves) and X onwards (peak table)."""
    from .sheets import select_sheet

    select_sheet(window, name)
    grid = window.peak_params_grid
    # Clear what an earlier save wrote (columns F..AP), merged cells included.
    last_col = 5 + 18 + len(PEAK_COLUMNS)
    for merged in list(ws.merged_cells.ranges):
        if merged.min_col <= last_col and merged.max_col >= 6:
            ws.unmerge_cells(str(merged))
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=6, max_col=last_col):
        for cell in row:
            cell.value = None
    if grid.GetNumberRows() == 0:
        return
    x = window.x_values
    fit, residuals = overall_fit(window)
    cols = [('BE', list(x))]
    if fit is not None:
        cols.append(('Residuals', list(residuals)))
    cols.append(('Background', list(window.background)))
    cols.append(('Calculated Fit', list(fit) if fit is not None else []))
    regions = None
    bg = window.Data['Core levels'][name].get('Background', {})
    if bg.get('Recorded_Ranges'):
        regions = np.zeros(len(x), dtype=bool)
        for r in bg['Recorded_Ranges']:
            if len(r) >= 4:
                lo, hi = sorted((float(r[2]), float(r[3])))
                regions |= (x >= lo) & (x <= hi)
    used = set(c for c, _ in cols)
    for i in range(grid.GetNumberRows() // 2):
        label = grid.GetCellValue(i * 2, 1)
        try:
            curve = peak_curve(window, i * 2, x)
        except (ValueError, TypeError, ZeroDivisionError):
            curve = None
        if curve is None:
            continue
        y = curve + window.background
        if regions is not None:
            y = np.where(regions, y, np.nan)
        unique = label
        k = 1
        while unique in used:
            unique = f"{label}_{k}"
            k += 1
        used.add(unique)
        cols.append((unique, _nan_to_none(y)))
    for j, (header, values) in enumerate(cols):
        c = 6 + j
        if c > 23:
            break
        ws.cell(row=1, column=c, value=header)
        for r, v in enumerate(values, start=2):
            ws.cell(row=r, column=c, value=None if v is None else float(v))
    for j, header in enumerate(PEAK_COLUMNS):
        ws.cell(row=1, column=24 + j, value=header)
        for r in range(grid.GetNumberRows()):
            v = grid.GetCellValue(r, j)
            ws.cell(row=2 + r, column=24 + j, value=v if v != '' else None)


def _write_results_table(wb, data):
    if 'Results Table' in wb.sheetnames:
        wb.remove(wb['Results Table'])
    ws = wb.create_sheet('Results Table')
    current_row = 2
    keys = sorted([k for k in data if k.startswith('Results Table')],
                  key=lambda x: int(x.replace('Results Table', '')) if x.replace('Results Table', '').isdigit() else 0)
    for key in keys:
        peaks = data[key].get('Peak') or {}
        if not peaks:
            continue
        ws.cell(row=current_row, column=2, value=f"Sample {key.replace('Results Table', '')}")
        current_row += 1
        for col, header in enumerate(RESULT_COLUMNS, start=2):
            ws.cell(row=current_row, column=col, value=header)
        ordered = sorted(((int(k[5:]), v) for k, v in peaks.items() if k.startswith('Peak_') and k[5:].isdigit()),
                         key=lambda t: t[0])
        for i, (_, p) in enumerate(ordered):
            for col_idx in range(len(RESULT_COLUMNS)):
                field = RESULT_FIELDS[col_idx] if col_idx < len(RESULT_FIELDS) else None
                value = p.get(field, "") if field else ""
                if col_idx == 7:
                    value = '✓' if value == '1' else ''
                if isinstance(value, (int, float)) and not isinstance(value, bool) and \
                        col_idx in [1, 2, 3, 4, 5, 6, 13, 14, 15]:
                    value = f"{value:.2f}"
                if value is None or isinstance(value, (dict, list)):
                    value = ""
                ws.cell(row=current_row + 1 + i, column=col_idx + 2, value=value)
        current_row = current_row + 1 + max(len(ordered), 1) - 1 + 4


def save_workbook(window, sheet_names, original_bytes=None):
    """Write the workbook and its JSON. Returns (xlsx bytes, json text)."""
    openpyxl = _openpyxl()
    data = window.Data
    if original_bytes:
        wb = openpyxl.load_workbook(io.BytesIO(original_bytes))
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    current = window.sheet_combobox.GetValue()
    # Writing each sheet selects it; the window's state comes back afterwards.
    keep = {k: getattr(window, k) for k in ('fit_results', 'vlines', 'selected_fitting_method', 'background_method',
                                            'offset_h', 'offset_l', 'bg_min_energy', 'bg_max_energy',
                                            'selected_peak_index')}
    for name in sheet_names:
        cl = data['Core levels'].get(name)
        if not cl or 'B.E.' not in cl:
            continue
        if name in wb.sheetnames:
            ws = wb[name]
        else:
            ws = wb.create_sheet(name[:31])
            ws.cell(row=1, column=1, value='Binding Energy')
            ws.cell(row=1, column=2, value='Raw Data')
            for r, (x, y) in enumerate(zip(cl['B.E.'], cl['Raw Data']), start=2):
                ws.cell(row=r, column=1, value=float(x))
                ws.cell(row=r, column=2, value=float(y))
        has_fit = isinstance(cl.get('Fitting'), dict) and cl['Fitting'].get('Peaks')
        if has_fit or ws.max_column > 5:
            _write_fit_columns(ws, window, name)
    # Results tables with no rows are dropped, as the desktop does.
    for key in [k for k in data if k.startswith('Results Table') and not (data[k].get('Peak'))]:
        del data[key]
    _write_results_table(wb, data)
    if current:
        from .sheets import select_sheet
        select_sheet(window, current)
        for k, v in keep.items():
            setattr(window, k, v)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), json.dumps(to_json_data(data), indent=2)
