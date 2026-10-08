# KherveOS: these definitions cut unchanged from KherveFittingPro origin/dev-AI (ca1fe50), libraries/ConfigFile.py. Regenerate with tools/export_khervetech.py.
import os
import sys
import json
import shutil
import wx


_APP_CONFIG_DIRNAME = 'KherveFitting'


def _app_root():
    """Folder the application runs from (this file lives in <root>/libraries)."""
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def get_config_dir():
    """Return (creating if needed) the per-user OS-standard config directory."""
    base = None
    try:
        base = wx.StandardPaths.Get().GetUserConfigDir()
    except Exception:
        base = None
    if not base:
        # Fallback if wx is unavailable for any reason.
        if sys.platform == 'win32':
            base = os.environ.get('APPDATA') or os.path.expanduser('~')
        elif sys.platform == 'darwin':
            base = os.path.join(os.path.expanduser('~'), 'Library', 'Preferences')
        else:
            base = os.environ.get('XDG_CONFIG_HOME') or \
                   os.path.join(os.path.expanduser('~'), '.config')
    path = os.path.join(base, _APP_CONFIG_DIRNAME)
    try:
        os.makedirs(path, exist_ok=True)
    except OSError:
        pass
    return path


def get_documents_dir():
    """The user's real Documents folder (handles OneDrive redirection)."""
    try:
        d = wx.StandardPaths.Get().GetDocumentsDir()
        if d:
            return d
    except Exception:
        pass
    return os.path.join(os.path.expanduser('~'), 'Documents')


def _ftir_library_seed_dir():
    """The read-only starter FTIR reference library shipped with the app."""
    if getattr(sys, 'frozen', False):
        from libraries.AppPaths import find_resource
        return find_resource('FTIR Library')
    return os.path.join(_app_root(), 'FTIR Library')


def get_ftir_library_dir():
    """Writable per-user folder for FTIR reference spectra (.jdx files).

    Seeded once from the bundled synthetic reference set; users can also save
    NIST WebBook downloads or their own JCAMP-DX files here."""
    dest = os.path.join(get_config_dir(), 'FTIR Library')
    marker = os.path.join(dest, '.seeded')
    try:
        os.makedirs(dest, exist_ok=True)
        if not os.path.exists(marker):
            src = _ftir_library_seed_dir()
            if src and os.path.isdir(src):
                existing = set(os.listdir(dest))
                for name in os.listdir(src):
                    if name in existing or name.startswith('_') or name.startswith('.'):
                        continue
                    s = os.path.join(src, name)
                    d = os.path.join(dest, name)
                    try:
                        if os.path.isdir(s):
                            shutil.copytree(s, d)
                        else:
                            shutil.copy2(s, d)
                    except OSError:
                        pass
            try:
                open(marker, 'w').close()
            except OSError:
                pass
    except OSError:
        pass
    return dest


def get_ftir_nicolet_dir():
    """Remembered folder of Nicolet/OMNIC (.lbd) libraries for the FTIR
    reference browser.

    These libraries are large (hundreds of MB) and shipped with the user's
    instrument, so they are browsed in place rather than copied into the app.
    The chosen folder is remembered in a tiny sidecar file (not the main
    config, so it can never clobber other preferences). Returns '' if unset or
    the folder no longer exists.
    """
    marker = os.path.join(get_config_dir(), 'ftir_nicolet_dir.json')
    try:
        with open(marker, 'r', encoding='utf-8') as f:
            folder = json.load(f).get('folder', '')
        if folder and os.path.isdir(folder):
            return folder
    except (OSError, ValueError):
        pass
    return ''


def set_ftir_nicolet_dir(folder):
    """Persist the FTIR Nicolet-library browser folder (see above)."""
    marker = os.path.join(get_config_dir(), 'ftir_nicolet_dir.json')
    try:
        with open(marker, 'w', encoding='utf-8') as f:
            json.dump({'folder': folder or ''}, f)
    except OSError:
        pass


FTIR_ONLINE_DEFAULT_BASE = (
    "https://raw.githubusercontent.com/gkerherve/Kherve-Downloads/"
    "main/FTIR_Reference/")


def get_ftir_online_base():
    """Base URL of the online (GitHub) FTIR reference library.

    The library is an ``index.json`` manifest plus per-spectrum ``.jdx`` files;
    only redistributable public-domain spectra (NIST / USGS / EPA) live there.
    Users can point this at their own fork/repo. Always ends with '/'.
    """
    marker = os.path.join(get_config_dir(), 'ftir_online_base.json')
    try:
        with open(marker, 'r', encoding='utf-8') as f:
            base = json.load(f).get('base', '')
        if base:
            return base if base.endswith('/') else base + '/'
    except (OSError, ValueError):
        pass
    return FTIR_ONLINE_DEFAULT_BASE


def set_ftir_online_base(base):
    """Persist the online FTIR reference library base URL (see above)."""
    marker = os.path.join(get_config_dir(), 'ftir_online_base.json')
    try:
        with open(marker, 'w', encoding='utf-8') as f:
            json.dump({'base': base or ''}, f)
    except OSError:
        pass


def get_ftir_online_cache_path():
    """Local cache file for the downloaded online-library index."""
    return os.path.join(get_config_dir(), 'ftir_online_index.json')


def get_xrd_library_dir():
    """Writable per-user folder for XRD crystal structures (.cif files).

    COD downloads and user CIFs land here; the XRD search-match scans it as
    the default local library."""
    dest = os.path.join(get_config_dir(), 'XRD Structures')
    try:
        os.makedirs(dest, exist_ok=True)
    except OSError:
        pass
    return dest


def get_xrd_instrument_path():
    """The per-user default XRD instrument parameter file (JSON)."""
    return os.path.join(get_config_dir(), 'xrd_instrument.json')


def Init_Measurement_Data(window):
    Data = {
        'FilePath': '',
        'Number of Core levels': 0,
        'Core levels': {},
    }

    # Initialize 10 Results Tables
    for i in range(10):
        Data[f'Results Table{i}'] = {
            'Peak': {}
        }

    return Data


def build_core_level_Data(sheet_name, be_values, plot_values, raw_values=None,
                          transmission=None, experimental_info=None):
    """In-memory counterpart of :func:`add_core_level_Data`, for importers.

    Builds the same core-level dictionary an importer would get by writing a
    ``BE | Corrected Data | Raw Data | Transmission`` sheet and reading it back,
    but without the workbook. Importers that can write a .kfit directly use
    this so both of their routes produce identical projects.

    Mind the reader's column mapping, reproduced here: column B — the
    transmission-corrected trace that actually gets plotted and fitted — lands
    in ``'Raw Data'``, and column C, the uncorrected counts, lands in
    ``'Corrected Data'``. ``plot_values`` is that column B trace;
    ``raw_values`` defaults to it (for sheets written without a separate raw
    column) and ``transmission`` to 1.0.

    Every column is cut to the shortest, as the sheet writers' ``zip()`` does —
    a B.E. axis longer than its counts crashes the plot.
    """
    be_values = list(be_values)
    plot_values = list(plot_values)
    raw_values = list(raw_values) if raw_values is not None else list(plot_values)
    transmission = (list(transmission) if transmission is not None
                    else [1.0] * len(plot_values))

    n = min(len(be_values), len(plot_values), len(raw_values), len(transmission))
    be_values = [float(v) for v in be_values[:n]]

    core_level = {
        'Name': sheet_name,
        'B.E.': be_values,
        'Raw Data': [float(v) for v in plot_values[:n]],
        'Corrected Data': [float(v) for v in raw_values[:n]],
        'Transmission': [float(v) for v in transmission[:n]],
    }
    if experimental_info:
        core_level['ExperimentalInfo'] = {
            str(k): ('' if v is None else str(v).strip())
            for k, v in experimental_info.items()
        }
    if be_values:
        core_level['Background'] = {
            'Bkg Y': list(core_level['Raw Data']),
            'Bkg Type': '',
            'Bkg Low': float(min(be_values)),
            'Bkg High': float(max(be_values)),
            'Bkg Offset Low': 0,
            'Bkg Offset High': 0,
        }
    return core_level


def add_core_level_Data(Data, window, file_path, sheet_name):
    """
    Add core level data from Excel sheet to Data structure, including experimental info
    """
    import openpyxl
    import pandas as pd
    import numpy as np

    try:
        # Read the Excel file
        df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)

        # ========== Handle EELS spectrum sheets ==========
        # 'EELS~Plot*' (legacy), 'EELS~LL', 'EELS~HL', pinned 'EELS~LL1', ...
        # — anything EELS that is not a map/image sheet.
        if sheet_name.startswith('EELS~') and not sheet_name.startswith('EELS~Map'):
            be_values = []
            raw_data = []

            for index, row in df.iterrows():
                if index == 0:
                    continue

                be_val = row.iloc[0]
                raw_val = row.iloc[1]

                if pd.isna(be_val) or pd.isna(raw_val):
                    continue

                try:
                    be_values.append(float(be_val))
                    raw_data.append(float(raw_val))
                except (ValueError, TypeError):
                    continue

            # Live extraction sheets are named after their signal
            signal = None
            suffix = sheet_name.split('~', 1)[-1]
            if suffix[:2] in ('LL', 'HL'):
                signal = suffix[:2]

            # Create complete Background structure with all required keys
            core_level = {
                'Name': sheet_name,
                'B.E.': be_values,
                'Raw Data': raw_data,
                '_EELS_type': 'plot',
                **({'_EELS_signal': signal} if signal else {}),
                'Background': {
                    'Bkg Type': '',
                    'Bkg Low': '',
                    'Bkg High': '',
                    'Bkg Offset Low': '',
                    'Bkg Offset High': '',
                    'Bkg X': be_values.copy() if be_values else [],
                    'Bkg Y': raw_data.copy() if raw_data else []
                }
            }

            Data['Core levels'][sheet_name] = core_level
            return Data

            # ========== Handle EELS map/image sheets ==========
        if sheet_name.startswith('EELS~Map'):
            # Rebuild from the header metadata + the preview rows. The .json
            # sibling normally supplies all of this; this path only runs when
            # the .xlsx is opened on its own.
            energy_range, signal, source, rotation = None, None, None, 0
            if len(df) > 0:
                for cell in df.iloc[0]:
                    text = str(cell)
                    if 'Range:' in text:
                        energy_range = text.split('Range:')[1].strip()
                    elif text.startswith('Signal:'):
                        signal = text.split(':', 1)[1].strip() or None
                    elif text.startswith('Source:'):
                        source = text.split(':', 1)[1].strip() or None
                    elif text.startswith('Rotation:'):
                        try:
                            rotation = int(text.split(':', 1)[1])
                        except (ValueError, TypeError):
                            pass

            # Preview rows start after the header + blank spacer row
            map_rows = []
            for index in range(2, len(df)):
                values = pd.to_numeric(df.iloc[index], errors='coerce')
                values = values.dropna().tolist()
                if not values:
                    break
                map_rows.append([float(v) for v in values])

            if not signal:
                # Legacy single-map sheet: hunt for the DM file next to the
                # project, as older versions of the importer relied on names.
                base_dir = os.path.dirname(file_path)
                stem = os.path.basename(file_path).replace('_EELS.xlsx', '')
                import glob
                for ext in ['.dm3', '.dm4', '.dm5']:
                    matches = ([p for p in [file_path.replace('.xlsx', ext),
                                            os.path.join(base_dir, stem + ext)]
                                if os.path.exists(p)]
                               or glob.glob(os.path.join(base_dir,
                                                         stem + '*' + ext)))
                    if matches:
                        source = os.path.basename(matches[0])
                        break

            core_level = {
                'Name': sheet_name,
                'Energy_Range': energy_range if energy_range else 'N/A',
                '_EELS_type': 'image' if signal == 'ADF' else 'map',
                '_EELS_signal': signal,
                '_EELS_source': source,
                '_EELS_rotation': rotation,
            }
            if map_rows:
                core_level['Map_Intensity'] = map_rows
                core_level['Map_Shape'] = [len(map_rows), len(map_rows[0])]

            Data['Core levels'][sheet_name] = core_level
            return Data

        # ========== Handle XPS~Map and other scienta-style map sheets ==========
        # Catches 'XPS~Map', 'C1s~Map', 'Zn2p~Map1', etc. (BE axis + Y1..Yn sweep columns).
        # EDX~Map / EELS~Map / ARPES~Map keep their own loaders and are excluded.
        if ('~Map' in sheet_name
                and not sheet_name.startswith(('EDX~Map', 'EELS~Map', 'ARPES~Map'))):
            be_values = []
            y_columns = {}
            num_sweeps = 0

            header_row = df.iloc[0] if len(df) > 0 else None
            if header_row is not None:
                for col_idx, val in enumerate(header_row):
                    if str(val).startswith('Y') and str(val)[1:].isdigit():
                        num_sweeps = max(num_sweeps, int(str(val)[1:]))

            for i in range(1, num_sweeps + 1):
                y_columns[f'Y{i}'] = []

            for index, row in df.iterrows():
                if index == 0:
                    continue
                be_val = row.iloc[0]
                if pd.isna(be_val):
                    continue
                try:
                    be_values.append(float(be_val))
                    for i in range(1, num_sweeps + 1):
                        col_idx = i
                        if col_idx < len(row) and not pd.isna(row.iloc[col_idx]):
                            y_columns[f'Y{i}'].append(float(row.iloc[col_idx]))
                        else:
                            y_columns[f'Y{i}'].append(0.0)
                except (ValueError, TypeError):
                    continue

            experimental_info = {}
            wb = openpyxl.load_workbook(file_path)
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                exp_col = None
                for col in range(num_sweeps + 5, min(num_sweeps + 30, ws.max_column + 1)):
                    cell_value = ws.cell(row=1, column=col).value
                    if cell_value and "Experimental Description" in str(cell_value):
                        exp_col = col
                        break
                if exp_col:
                    for row in range(2, ws.max_row + 1):
                        param_cell = ws.cell(row=row, column=exp_col)
                        value_cell = ws.cell(row=row, column=exp_col + 1)
                        if param_cell.value is not None and str(param_cell.value).strip():
                            experimental_info[str(param_cell.value).strip()] = str(value_cell.value).strip() if value_cell.value else ""
            wb.close()

            core_level = {
                'Name': sheet_name,
                'B.E.': [float(f"{val:.2f}") for val in be_values],
                '_Map_type': 'scienta',
                '_num_sweeps': num_sweeps,
                '_core_level': experimental_info.get('Core Level', ''),
            }
            for col_name, values in y_columns.items():
                core_level[col_name] = [float(f"{val:.2f}") for val in values]
            if experimental_info:
                core_level['ExperimentalInfo'] = experimental_info

            # 2D spatial (XY) area maps carry a pixel-grid calibration so they can be
            # shown as an X/Y image and summed over any BE window.
            if str(experimental_info.get('Map Kind', '')).upper() == 'XY':
                try:
                    core_level['_Map_type'] = 'vgd_xy'
                    core_level['_n_x'] = int(float(experimental_info.get('X Pixels', 0)))
                    core_level['_n_y'] = int(float(experimental_info.get('Y Pixels', 0)))
                    core_level['_x_step'] = float(experimental_info.get('X Step (um)', 0) or 0)
                    core_level['_y_step'] = float(experimental_info.get('Y Step (um)', 0) or 0)
                except (ValueError, TypeError):
                    pass
            if be_values:
                core_level['Background'] = {
                    'Bkg Type': '',
                    'Bkg Low': float(f"{min(be_values):.2f}"),
                    'Bkg High': float(f"{max(be_values):.2f}"),
                    'Bkg Offset Low': 0,
                    'Bkg Offset High': 0
                }
            Data['Core levels'][sheet_name] = core_level
            return Data

        # Extract B.E. and Raw Data columns
        be_values = []
        raw_data = []
        corrected_data = []
        transmission = []

        # Skip the header row and extract data
        for index, row in df.iterrows():
            if index == 0:  # Skip header
                continue

            be_val = row.iloc[0]  # Column A (B.E.)
            raw_val = row.iloc[1]  # Column B (Raw Data)

            # Handle missing or invalid data
            if pd.isna(be_val) or pd.isna(raw_val):
                continue

            try:
                be_values.append(float(be_val))
                raw_data.append(float(raw_val))

                # Check if corrected data column exists
                if len(row) > 2 and not pd.isna(row.iloc[2]):
                    corrected_data.append(float(row.iloc[2]))
                else:
                    corrected_data.append(float(raw_val))

                # Check if transmission column exists
                if len(row) > 3 and not pd.isna(row.iloc[3]):
                    transmission.append(float(row.iloc[3]))
                else:
                    transmission.append(1.0)

            except (ValueError, TypeError):
                continue

        # Extract experimental description data from Excel file
        experimental_info = {}

        # Load workbook to access experimental description columns
        wb = openpyxl.load_workbook(file_path)
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]

            # Search for experimental description column (typically around column 45-50)
            exp_col = None
            for col in range(40, min(61, ws.max_column + 1)):
                cell_value = ws.cell(row=1, column=col).value
                if cell_value and "Experimental Description" in str(cell_value):
                    exp_col = col
                    break

            if exp_col:
                # Read experimental description data
                for row in range(2, ws.max_row + 1):
                    param_cell = ws.cell(row=row, column=exp_col)
                    value_cell = ws.cell(row=row, column=exp_col + 1)

                    if param_cell.value is not None and str(param_cell.value).strip():
                        param_name = str(param_cell.value).strip()
                        param_value = str(value_cell.value).strip() if value_cell.value is not None else ""
                        experimental_info[param_name] = param_value
        # Windows keeps the .xlsx locked while the workbook is open, which is
        # what stopped an importer's throwaway intermediate from being deleted
        # once the project was converted to .kfit.
        wb.close()

        # Create the core level data structure with .2f formatting
        core_level_data = {
            'B.E.': [float(f"{val:.2f}") for val in be_values],
            'Raw Data': [float(f"{val:.2f}") for val in raw_data],
            'Corrected Data': [float(f"{val:.2f}") for val in corrected_data],
            'Transmission': [float(f"{val:.2f}") for val in transmission],
            'Name': sheet_name
        }

        # Add experimental info to the core level data if found
        if experimental_info:
            core_level_data['ExperimentalInfo'] = experimental_info

        # Initialize background structure
        if be_values:
            core_level_data['Background'] = {
                'Bkg Y': core_level_data['Raw Data'],
                'Bkg Type': '',
                'Bkg Low': float(f"{min(be_values):.2f}"),
                'Bkg High': float(f"{max(be_values):.2f}"),
                'Bkg Offset Low': 0,
                'Bkg Offset High': 0
            }

        # Add to Data structure
        if 'Core levels' not in Data:
            Data['Core levels'] = {}

        Data['Core levels'][sheet_name] = core_level_data
        Data['Number of Core levels'] = len(Data['Core levels'])

        return Data

    except Exception as e:
        print(f"Error adding core level data for {sheet_name}: {e}")
        return Data


def add_peak_to_core_level_Data(data, core_name, peak_data):
    if core_name in data['Core levels']:
        fitting = data['Core levels'][core_name]['Fitting']
        fitting.update(peak_data)
    else:
        print(f"Core level {core_name} does not exist.")

